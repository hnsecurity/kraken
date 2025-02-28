from tqdm import tqdm
import re
from openpyxl.chart import PieChart3D, PieChart, ProjectedPieChart, BarChart, Reference, Series
from openpyxl.styles import Font
import operator
from Levenshtein import distance
import logging

from google_api import GoogleApi
import pwnedpasswords

logger = logging.getLogger('logger')


def password_analysis_stats(ws_stats, start_line, args, list_users, list_user_password):
    ws_stats[f"B{start_line}"] = "Passwords Analysis"
    ws_stats[f"B{start_line}"].font = Font(bold=True)

    ws_stats[f"B{start_line+1}"] = "Analyzed"
    ws_stats[f"C{start_line+1}"] = len(list_user_password)

    ws_stats[f"B{start_line+2}"] = "Not analyzed*"

    if list_users:
        ws_stats[f"C{start_line+2}"] = len(list_users) - \
            len(list_user_password)
    else:
        ws_stats[f"C{start_line+2}"] = 0

    chart = PieChart3D()
    chart.add_data(
        Reference(ws_stats, range_string=f'Stats!C{start_line+1}:C{start_line+2}'))
    chart.set_categories(
        Reference(ws_stats, range_string=f'Stats!B{start_line+1}:B{start_line+2}'))
    chart.title = "Password Analysis"
    ws_stats.add_chart(chart, f'E{start_line}')

    return start_line + 15


def characters_analysis_stats(ws_stats, start_line, args, list_users, list_user_password):
    password_chars = {}
    password_chars["Uppercase"] = 0
    password_chars["Lowercase"] = 0
    password_chars["Digits"] = 0
    password_chars["Specials"] = 0
    password_chars["Digits Lowercase"] = 0
    password_chars["Digits Uppercase"] = 0
    password_chars["Digits Specials"] = 0
    password_chars["Lowercase Uppercase"] = 0
    password_chars["Lowercase Specials"] = 0
    password_chars["Uppercase Specials"] = 0
    password_chars["Digits Lowercase Uppercase"] = 0
    password_chars["Digits Lowercase Specials"] = 0
    password_chars["Digits Uppercase Specials"] = 0
    password_chars["Lowercase Uppercase Specials"] = 0
    password_chars["Digits Lowercase Uppercase Specials"] = 0

    for cur_user in list_user_password:

        # don't count on password history
        if re.search("_history[0-9]*$", cur_user):
            continue

        # check password chars
        if re.search("(?=^[^\x0d\x0a]*$)(?!.*\d)(?=.*[A-Z])(?!.*[a-z])(?!.*[\W])", list_user_password[cur_user]):
            password_chars["Uppercase"] += 1
            continue
        if re.search("(?=^[^\x0d\x0a]*$)(?!.*\d)(?!.*[A-Z])(?=.*[a-z])(?!.*[\W])", list_user_password[cur_user]):
            password_chars["Lowercase"] += 1
            continue
        if re.search("(?=^[^\x0d\x0a]*$)(?=.*\d)(?!.*[A-Z])(?!.*[a-z])(?!.*[\W])", list_user_password[cur_user]):
            password_chars["Digits"] += 1
            continue
        if re.search("(?=^[^\x0d\x0a]*$)(?!.*\d)(?!.*[A-Z])(?!.*[a-z])(?=.*[\W])", list_user_password[cur_user]):
            password_chars["Specials"] += 1
            continue
        if re.search("(?=^[^\x0d\x0a]*$)(?=.*\d)(?!.*[A-Z])(?=.*[a-z])(?!.*[\W])", list_user_password[cur_user]):
            password_chars["Digits Lowercase"] += 1
            continue
        if re.search("(?=^[^\x0d\x0a]*$)(?=.*\d)(?=.*[A-Z])(?!.*[a-z])(?!.*[\W])", list_user_password[cur_user]):
            password_chars["Digits Uppercase"] += 1
            continue
        if re.search("(?=^[^\x0d\x0a]*$)(?=.*\d)(?!.*[A-Z])(?!.*[a-z])(?=.*[\W])", list_user_password[cur_user]):
            password_chars["Digits Specials"] += 1
            continue
        if re.search("(?=^[^\x0d\x0a]*$)(?!.*\d)(?=.*[A-Z])(?=.*[a-z])(?!.*[\W])", list_user_password[cur_user]):
            password_chars["Lowercase Uppercase"] += 1
            continue
        if re.search("(?=^[^\x0d\x0a]*$)(?!.*\d)(?!.*[A-Z])(?=.*[a-z])((?=.*[\W])|(?=.*[\s]))", list_user_password[cur_user]):
            password_chars["Lowercase Specials"] += 1
            continue
        if re.search("(?=^[^\x0d\x0a]*$)(?!.*\d)(?=.*[A-Z])(?!.*[a-z])(?=.*[\W])", list_user_password[cur_user]):
            password_chars["Uppercase Specials"] += 1
            continue
        if re.search("(?=^[^\x0d\x0a]*$)(?=.*\d)(?=.*[A-Z])(?=.*[a-z])(?!.*[\W])", list_user_password[cur_user]):
            password_chars["Digits Lowercase Uppercase"] += 1
            continue
        if re.search("(?=^[^\x0d\x0a]*$)(?=.*\d)(?!.*[A-Z])(?=.*[a-z])(?=.*[\W])", list_user_password[cur_user]):
            password_chars["Digits Lowercase Specials"] += 1
            continue
        if re.search("(?=^[^\x0d\x0a]*$)(?=.*\d)(?=.*[A-Z])(?!.*[a-z])(?=.*[\W])", list_user_password[cur_user]):
            password_chars["Digits Uppercase Specials"] += 1
            continue
        if re.search("(?=^[^\x0d\x0a]*$)(?!.*\d)(?=.*[A-Z])(?=.*[a-z])(?=.*[\W])", list_user_password[cur_user]):
            password_chars["Lowercase Uppercase Specials"] += 1
            continue
        if re.search("(?=^[^\x0d\x0a]*$)(?=.*\d)(?=.*[A-Z])(?=.*[a-z])(?=.*[\W])", list_user_password[cur_user]):
            password_chars["Digits Lowercase Uppercase Specials"] += 1
            continue

    ws_stats[f"B{start_line}"] = "Characters Analysis"
    ws_stats[f"B{start_line}"].font = Font(bold=True)

    cur_line = start_line + 1

    if args.show_not_cracked and list_users:
        ws_stats["B" + str(cur_line)] = "Not cracked"
        ws_stats["C" + str(cur_line)] = len(list_users) - \
            len(list_user_password)
        cur_line += 1

    for i in sorted(password_chars.items(), key=operator.itemgetter(1), reverse=True):
        if i[1] > 0:
            ws_stats["B" + str(cur_line)] = i[0]
            ws_stats["C" + str(cur_line)] = i[1]
            cur_line += 1

    chart = ProjectedPieChart()
    chart.add_data(
        Reference(ws_stats, range_string=f'Stats!C{start_line+1}:C{cur_line-1}'))
    chart.set_categories(
        Reference(ws_stats, range_string=f'Stats!B{start_line+1}:B{cur_line-1}'))
    chart.title = "Characters Analysis"
    ws_stats.add_chart(chart, f'E{start_line}')

    return max(cur_line, start_line+15)


def password_length_stats(ws_stats, start_line, args, list_users, list_user_password):
    cut_length = 14

    password_len = {}
    for i in range(0, cut_length+2):
        password_len[i] = 0

    for cur_user in list_user_password:
        # compute len of password
        if len(list_user_password[cur_user]) > cut_length:
            password_len[cut_length+1] += 1
        else:
            password_len[len(list_user_password[cur_user])] += 1

    ws_stats[f"B{start_line}"] = "Password Length"
    ws_stats[f"B{start_line}"].font = Font(bold=True)

    cur_line = start_line + 1

    if args.show_not_cracked and list_users:
        ws_stats["B" + str(cur_line)] = "NC"
        ws_stats["C" + str(cur_line)] = len(list_users) - \
            len(list_user_password)
        cur_line += 1

    max_length = max([i for i in range(cut_length+1) if password_len[i] != 0]+[0])
    if password_len[cut_length+1] != 0:
        max_length = cut_length
    
    for i in range(0, max_length+1):
        ws_stats["B" + str(cur_line)] = i
        ws_stats["C" + str(cur_line)] = password_len[i]
        cur_line += 1

    if password_len[cut_length+1] != 0:
        ws_stats["B" + str(cur_line)] = f">{cut_length}"
        ws_stats["C" + str(cur_line)] = password_len[cut_length+1]
        cur_line += 1

    chart = BarChart()
    chart.add_data(
        Reference(ws_stats, range_string=f'Stats!C{start_line+1}:C{cur_line-1}'))
    chart.set_categories(
        Reference(ws_stats, range_string=f'Stats!B{start_line+1}:B{cur_line-1}'))
    chart.title = "Password length"
    chart.legend = None
    ws_stats.add_chart(chart, f'E{start_line}')

    return max(cur_line, start_line+15)


def convert_into_l33t_regex(s):
    n = ""

    for c in s:
        n = n + '[' + str(c) + c.swapcase()

        if c == 'a':
            n = n + '4'
        elif c == 'e':
            n = n + '3'
        elif c == 'i':
            n = n + '1'
        elif c == 'o':
            n = n + '0'

        n = n + ']+'

    return n

def password_leaked_stats(ws_stats, start_line, args, list_users, list_user_password, list_regex, dictionary_list):
    leaked_passwords_google = 0
    leaked_passwords_HIBP = 0
    count_leaked = 0
    count_not_leaked = 0

    if args.check_leaked_google:
        logger.debug('Getting token to use the google api (leaked passwords)')
        google_leak_api = GoogleApi()
        logger.debug('Google token found')

    leaked_set_HIBP = set()
    leaked_set_google = set()

    for cur_user in tqdm(list_user_password):
        # don't count on password history
        if re.search("_history[0-9]*$", cur_user):
            continue
        
        leaked = False
        # check if password was leaked

        if args.check_leaked_google:
            # clean the username
            cleaned_username = cur_user.split('\\')[-1]
            if google_leak_api.is_leaked(cleaned_username, list_user_password[cur_user]):
                leaked_passwords_google += 1
                leaked_set_google.add((cleaned_username,list_user_password[cur_user]))
                leaked = True

        if pwnedpasswords.check(list_user_password[cur_user]):
            leaked_passwords_HIBP += 1
            leaked_set_HIBP.add(list_user_password[cur_user])
            leaked = True

        if not leaked:
            count_not_leaked += 1
        else:
            count_leaked += 1

    logger.debug('Leaked passwords HIBP: ' + str(leaked_set_HIBP))
    logger.debug('Leaked passwords Google: ' + str(leaked_set_google))

    ws_stats[f"B{start_line}"] = "Leaked Password"
    ws_stats[f"B{start_line}"].font = Font(bold=True)

    cur_line = start_line+1

    ws_stats["B" + str(cur_line)] = 'Leaked'
    ws_stats["C" + str(cur_line)] = count_leaked
    cur_line += 1

    ws_stats["B" + str(cur_line)] = 'Not leaked'
    ws_stats["C" + str(cur_line)] = count_not_leaked
    cur_line += 1

    if args.check_leaked_google:
        ws_stats["B" + str(cur_line)] = 'Leaked HIBP'
        ws_stats["C" + str(cur_line)] = leaked_passwords_HIBP
        cur_line += 1

        ws_stats["B" + str(cur_line)] = 'Leaked Google'
        ws_stats["C" + str(cur_line)] = leaked_passwords_google
        cur_line += 1

    chart = BarChart()
    chart.add_data(
        Reference(ws_stats, range_string=f'Stats!C{start_line+1}:C{cur_line-1}'))
    chart.set_categories(
        Reference(ws_stats, range_string=f'Stats!B{start_line+1}:B{cur_line-1}'))
    chart.title = "Leaked Password"
    chart.legend = None
    ws_stats.add_chart(chart, f'E{start_line}')

    return max(cur_line, start_line+15)

def password_topology_stats(ws_stats, start_line, args, list_users, list_user_password, list_regex, dictionary_list):
    username_based = 0
    username_same = 0
    repeated_word = 0

    regex_results = {}
    dictionary_results = {}

    for cur_user in list_user_password:

        # don't count on password history
        if re.search("_history[0-9]*$", cur_user):
            continue

        # check if password is the username
        if cur_user.upper() == list_user_password[cur_user].upper():
            username_same += 1
        else:
            # check if password is derived from username
            r = convert_into_l33t_regex(cur_user)
            r_i = convert_into_l33t_regex(cur_user[::-1])

            if re.search(r, list_user_password[cur_user]) or re.search(r_i, list_user_password[cur_user]):
                username_based += 1

        # check is in the password there is repeated words
        for i in range(0, len(list_user_password[cur_user]) - 2):
            sub = re.escape(list_user_password[cur_user][i:i+3])

            if len(re.findall(sub, list_user_password[cur_user])) > 1:

                repeated_word += 1
                break

        # check password on regex
        try:
            for reg in list_regex:
                if re.search(reg, list_user_password[cur_user]):
                    try:
                        regex_results[reg] += 1
                    except:
                        regex_results[reg] = 1
        except:
            pass

        # check password on dictionary
        try:
            for w in dictionary_list:
                r = convert_into_l33t_regex(w)
                if re.search(r, list_user_password[cur_user]):
                    if w in dictionary_results:
                        dictionary_results[w] += 1
                    else:
                        dictionary_results[w] = 1
        except:
            pass
    # create password type stats
    password_type = {}

    password_type["Same as username"] = username_same
    password_type["Based on username"] = username_based
    password_type["Repeated word"] = repeated_word
    password_type["Others"] = len(list_user_password) - username_same - username_based - repeated_word


    try:
        for i in list_regex:
            try:
                password_type["Based on \"" + str(i) + "\""] = regex_results[i]
            except:
                pass
    except:
        pass

    try:
        for i in dictionary_results:
            try:
                password_type["Based on \"" +
                              str(i) + "\""] = dictionary_results[i]
            except:
                pass
    except:
        pass

    ws_stats[f"B{start_line}"] = "Password Topology"
    ws_stats[f"B{start_line}"].font = Font(bold=True)

    cur_line = start_line+1

    for i in sorted(password_type.items(), key=operator.itemgetter(1), reverse=True):
        if i[1] > 0:
            ws_stats["B" + str(cur_line)] = i[0]
            ws_stats["C" + str(cur_line)] = i[1]
            cur_line += 1

    chart = BarChart()
    chart.add_data(
        Reference(ws_stats, range_string=f'Stats!C{start_line+1}:C{cur_line-1}'))
    chart.set_categories(
        Reference(ws_stats, range_string=f'Stats!B{start_line+1}:B{cur_line-1}'))
    chart.title = "Password Topology"
    chart.legend = None
    ws_stats.add_chart(chart, f'E{start_line}')

    return max(cur_line, start_line+15)


def levenshtein_distance_stats(ws_stats, start_line, args, list_users, list_user_password, list_user_password_with_history):
    
    cut_distance = 14
    
    levenshtein_distance = {}
    for i in range(0, cut_distance+2):
        levenshtein_distance[i] = 0

    for cur_user in list_user_password:

        # don't count on password history
        if re.search("_history[0-9]*$", cur_user):
            continue

        # check Levenshtein distance if history0 exist
        try:

            if list_user_password_with_history[cur_user + "_history0"]:

                if args.ignore_case:
                    dis = distance(list_user_password[cur_user].upper(), list_user_password_with_history[cur_user + "_history0"].upper())
                else:
                    dis = distance(list_user_password[cur_user], list_user_password_with_history[cur_user + "_history0"])

                if dis > cut_distance:
                    levenshtein_distance[cut_distance+1] += 1
                else:
                    levenshtein_distance[dis] += 1

        except:
            pass

    


    ws_stats[f"B{start_line}"] = "Levenshtein distance"
    ws_stats[f"B{start_line}"].font = Font(bold=True)

    cur_line = start_line + 1

    if args.show_not_cracked and list_users:
        ws_stats["B" + str(cur_line)] = "NC"
        ws_stats["C" + str(cur_line)] = len(list_users) - \
            len(list_user_password)
        cur_line += 1

        ws_stats["B" + str(cur_line)] = "Unk"
        ws_stats["C" + str(cur_line)] = len(list_user_password) - \
            sum(levenshtein_distance.values())
        cur_line += 1

    max_dist = max([i for i in range(cut_distance+1) if levenshtein_distance[i] != 0]+[0])

    if levenshtein_distance[cut_distance+1] != 0:
        max_dist = cut_distance

    for i in range(max_dist+1):
        ws_stats["B" + str(cur_line)] = i
        ws_stats["C" + str(cur_line)] = levenshtein_distance[i]
        cur_line += 1

    if levenshtein_distance[cut_distance+1] != 0:
        ws_stats["B" + str(cur_line)] = f">{cut_distance}"
        ws_stats["C" + str(cur_line)] = levenshtein_distance[cut_distance+1]
        cur_line += 1

    chart = BarChart()
    chart.add_data(
        Reference(ws_stats, range_string=f'Stats!C{start_line+1}:C{cur_line-1}'))
    chart.set_categories(
        Reference(ws_stats, range_string=f'Stats!B{start_line+1}:B{cur_line-1}'))
    chart.title = "Levenshtein distance"
    chart.legend = None
    ws_stats.add_chart(chart, f'E{start_line}')

    return max(cur_line, start_line+15)
